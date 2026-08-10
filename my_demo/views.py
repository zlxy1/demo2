import io
from io import BytesIO
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from openpyxl.reader.excel import load_workbook
from openpyxl.workbook import Workbook
from . import models
from .models import Staff, work_status_choices, area_choice, staff_type_choice
from django.core.paginator import Paginator
from django.contrib import messages
from account.views import admin_required
from account.models import UserProfile,DEP_CHOICE

def excel_export(data_list):
    wb=Workbook()
    ws=wb.active
    ws.title='员工信息表'
    row=['姓名','性别','出生日期','身份证号','人员类型','在职状态','所属科室']
    ws.append(row)
    for item in data_list:
        ws.append([item.name,item.get_gender_display(),item.birth,item.id_number,item.staff_type,item.work_status,item.work_area])
    buffer=io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def excel_import(file_bytes):
    wb=load_workbook(BytesIO(file_bytes),read_only=True)
    ws=wb.active
    datalist=[]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not any(row):
            continue
        name,gender,birth,id_number,staff_type,work_status,work_area=row
        qs=Staff(name=name,gender=gender,birth=birth,id_number=id_number,staff_type=staff_type,work_status=work_status,work_area=work_area)
        datalist.append(qs)
    Staff.objects.bulk_create(datalist)
    wb.close()
    msg=f'导入成功，共导入{len(datalist)}条数据'
    return msg

def staff_filter(request,qs):
    search_name=request.GET.get('name','')
    search_work_area=request.GET.get('work_area','')
    search_staff_type=request.GET.get('staff_type','')
    if search_name:
        qs = qs.filter(name__icontains=search_name)
    if search_work_area:
        qs = qs.filter(work_area=search_work_area)
    if search_staff_type:
        qs=qs.filter(staff_type=search_staff_type)
    return qs

@login_required
@admin_required
def staff_add(request):
    if request.method == 'GET':
        ctx = {'gender_list': models.gender_choice,
               'staff_type_list': models.staff_type_choice,
               'work_status_list': work_status_choices,
               'area_choice_list': area_choice
               }
        return render(request, 'staff/add.html', ctx)
    if request.method == 'POST':
        Staff.objects.create(name=request.POST.get('name'),
                             gender=request.POST.get('gender'),
                             birth=request.POST.get('birth'),
                             id_number=request.POST.get('id_number'),
                             staff_type=request.POST.get('staff_type'),
                             work_status=request.POST.get('work_status'),
                             work_area=request.POST.get('work_area'))
        return redirect('staff_list')

@login_required
@admin_required
def staff_del(request, pk):
    staff_del = get_object_or_404(Staff, pk=pk)
    if request.method == "POST":
        staff_del.delete()
        return redirect('staff_list')

@login_required
@admin_required
def staff_edit(request, pk):
    staff = get_object_or_404(Staff, pk=pk)
    ctx1 = {'gender_list': models.gender_choice,
            'staff_type_list': models.staff_type_choice,
            'work_status_list': work_status_choices,
            'area_choice_list': area_choice,
            'staff': staff
            }
    if request.method == 'GET':
        return render(request, 'staff/edit.html', ctx1)
    else:
        staff.name=request.POST.get('name')
        staff.gender = request.POST.get('gender')
        staff.birth = request.POST.get('birth')
        staff.id_number = request.POST.get('id_number')
        staff.staff_type = request.POST.get('staff_type')
        staff.work_status = request.POST.get('work_status')
        staff.save()
        return redirect('staff_list')

@login_required
@login_required
def staff_list(request):
    qs=Staff.objects.all().order_by('id')
    qs=staff_filter(request,qs)

    profile, _ = UserProfile.objects.get_or_create(user=request.user)
    is_admin = profile.is_admin or request.user.is_superuser
    user_dep = profile.dep

    if not is_admin:
        qs = qs.filter(work_area=user_dep)

    paginator = Paginator(qs, 10)
    pag_num = request.GET.get('page', 1)
    page_data = paginator.get_page(pag_num)

    return render(request,'staff/list.html',{
        'staff_list':page_data,
        'staff_type':staff_type_choice,
        'work_area':area_choice,
        'is_admin': is_admin,
        'user_dep': user_dep   # 新增传给模板
    })

@login_required
@admin_required
def staff_export(request):
    qs = Staff.objects.all().order_by('id')
    qs = staff_filter(request, qs)
    buf=excel_export(qs)
    response=HttpResponse(buf.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="员工信息表.xlsx"'
    return response

@login_required
@admin_required
def staff_import(request):
    if request.method=="POST":
        file=request.FILES.get('excel')
        if not file:
            return render(request,'staff/import.html',{'msg':'请选择要上传的数据表文件'})
        file_bytes=file.read()
        msg=excel_import(file_bytes)
        return render(request,'staff/import.html',{'msg':msg})
    return render(request,'staff/import.html')


@login_required
@admin_required
def ids_delete(request):
    if request.method == "POST":
        ids_str = request.POST.get('ids', '').strip()
        # 1. 判断是否为空
        if not ids_str:
            messages.error(request, "未勾选任何数据！")
            return redirect('staff_list')

        # 2. 分割并过滤空字符串（防止出现 ,, 产生空白元素）
        ids_list = [i.strip() for i in ids_str.split(',') if i.strip()]

        # 3. 防止过滤后数组为空
        if not ids_list:
            messages.error(request, "有效数据ID为空！")
            return redirect('staff_list')

        # 批量删除
        Staff.objects.filter(id__in=ids_list).delete()
        messages.success(request, "选中记录删除成功！")

    return redirect('staff_list')


